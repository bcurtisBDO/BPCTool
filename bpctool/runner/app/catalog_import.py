"""Build catalog work-item drafts from Excel with HTML/markdown normalization."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .excel_ir import _cell, _looks_like_catalog
from .mapping import TemplateMapping, normalize
from .richtext import normalize_field_value, normalize_rich_text

WORK_ITEM_TYPE_ALIASES = {
    "test case": "Test Case",
    "testcase": "Test Case",
    "process area": "Process area",
    "functional area": "Functional area",
    "system process": "System process",
    "end to end": "End to end",
}


@dataclass
class WorkItemDraft:
    key: str
    parent_key: str | None
    work_item_type: str
    title: str
    fields: dict[str, Any]
    source_file: str
    source_row: int
    content_formats: dict[str, str] = field(default_factory=dict)
    # Field refs that should be created with /multilineFieldsFormat = Markdown
    markdown_format_fields: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def load_catalog_drafts(
    catalog_dir: str | Path | None,
    mapping: TemplateMapping,
    *,
    limit: int | None = None,
    convert_markdown_to_html: bool = False,
) -> list[WorkItemDraft]:
    if not catalog_dir:
        return []
    root = Path(catalog_dir).expanduser().resolve()
    if not root.exists():
        raise FileNotFoundError(f"Catalog source folder not found: {root}")

    drafts: list[WorkItemDraft] = []
    for path in sorted(root.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        if not _looks_like_catalog(path) and "jul" not in path.name.lower():
            # Accept all xlsx in the BPC Files folder.
            pass
        drafts.extend(
            _drafts_from_workbook(
                path, mapping, convert_markdown_to_html=convert_markdown_to_html
            )
        )
        if limit is not None and len(drafts) >= limit:
            return drafts[:limit]
    return sort_parent_first(drafts)


def _drafts_from_workbook(
    path: Path,
    mapping: TemplateMapping,
    *,
    convert_markdown_to_html: bool = False,
) -> list[WorkItemDraft]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []

    header_idx = 0
    if len(rows) > 1 and any(_cell(v).lower() == "work item type" for v in rows[1]):
        header_idx = 1
    headers = [_cell(v) for v in rows[header_idx]]
    drafts: list[WorkItemDraft] = []
    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        values = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        if not any(v not in (None, "") for v in values.values()):
            continue
        values["_source_file"] = path.name
        values["_source_sheet"] = sheet_name
        values["_source_row"] = offset
        drafts.append(
            _row_to_draft(
                values, mapping, convert_markdown_to_html=convert_markdown_to_html
            )
        )
    return drafts


def _row_to_draft(
    row: dict[str, Any],
    mapping: TemplateMapping,
    *,
    convert_markdown_to_html: bool = False,
) -> WorkItemDraft:
    lookup = _build_lookup(row)
    wit = _work_item_type(lookup, row)
    hierarchy_title, hierarchy_level = _title_from_row(lookup)
    title = hierarchy_title or _first_value(
        lookup,
        "title",
        "name",
        "work item title",
        "process name",
        "business process",
        "deliverable",
        "test case",
        "description",
    )
    key = _external_key(lookup, row)
    if not title:
        title = f"{wit} {key}"

    fields: dict[str, Any] = {"System.Title": title}
    formats: dict[str, str] = {}
    markdown_format_fields: list[str] = []
    description = _first_value(lookup, "description")
    if description:
        is_html = _looks_like_html_marker(description)
        fields["System.Description"] = normalize_rich_text(
            description,
            field_type="html",
            convert_markdown_to_html=convert_markdown_to_html,
        )
        if is_html:
            formats["System.Description"] = "html"
        elif convert_markdown_to_html:
            formats["System.Description"] = "markdown→html"
        else:
            formats["System.Description"] = "markdown"
            markdown_format_fields.append("System.Description")
    else:
        fields["System.Description"] = (
            f"Imported from {row.get('_source_file')} row {row.get('_source_row')}."
        )
        formats["System.Description"] = "generated"
        if not convert_markdown_to_html:
            markdown_format_fields.append("System.Description")

    for header, value in row.items():
        if str(header).startswith("_") or value in (None, ""):
            continue
        ref = mapping.field_ref_for_header(str(header))
        if not ref or ref in {"System.Title", "System.Description"}:
            continue
        fd = mapping.field_def(ref)
        field_type = fd.field_type if fd else ("html" if mapping.is_html_field(ref) else "")
        if mapping.is_html_field(ref) or "html" in field_type.lower():
            is_html = _looks_like_html_marker(value)
            fields[ref] = normalize_field_value(
                value,
                field_type or "html",
                convert_markdown_to_html=convert_markdown_to_html,
            )
            if is_html:
                formats[ref] = "html"
            elif convert_markdown_to_html:
                formats[ref] = "markdown→html"
            else:
                formats[ref] = "markdown"
                markdown_format_fields.append(ref)
        else:
            fields[ref] = normalize_field_value(
                value, field_type, convert_markdown_to_html=convert_markdown_to_html
            )

    fields.setdefault("MSBPC.microsoftid", _first_value(lookup, "microsoft id") or key)
    fields.setdefault("MSBPC.processsequenceid", _first_value(lookup, "process sequence id") or key)

    return WorkItemDraft(
        key=key,
        parent_key=_parent_key(lookup, hierarchy_level),
        work_item_type=wit,
        title=title,
        fields={ref: value for ref, value in fields.items() if value not in (None, "")},
        source_file=str(row.get("_source_file") or ""),
        source_row=int(row.get("_source_row") or 0),
        content_formats=formats,
        markdown_format_fields=markdown_format_fields,
    )


def _looks_like_html_marker(value: Any) -> bool:
    from .richtext import looks_like_html

    return looks_like_html(value)


def sort_parent_first(drafts: list[WorkItemDraft]) -> list[WorkItemDraft]:
    by_key = {d.key: d for d in drafts}
    ordered: list[WorkItemDraft] = []
    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(draft: WorkItemDraft) -> None:
        if draft.key in visited:
            return
        if draft.key in visiting:
            draft.warnings.append("Cycle detected in parent relationship.")
            return
        visiting.add(draft.key)
        if draft.parent_key and draft.parent_key in by_key:
            visit(by_key[draft.parent_key])
        visiting.remove(draft.key)
        visited.add(draft.key)
        ordered.append(draft)

    for draft in drafts:
        visit(draft)
    return ordered


def _build_lookup(row: dict[str, Any]) -> dict[str, Any]:
    lookup: dict[str, Any] = {}
    for header, value in row.items():
        key = normalize(header)
        lookup[key] = value
    return lookup


def _title_from_row(lookup: dict[str, Any]) -> tuple[str | None, int | None]:
    titles: list[tuple[int, str]] = []
    for index in range(1, 11):
        value = _first_value(lookup, f"title {index}")
        if value:
            titles.append((index, value))
    if titles:
        return titles[-1][1], titles[-1][0]
    return None, None


def _work_item_type(lookup: dict[str, Any], row: dict[str, Any]) -> str:
    raw = _first_value(lookup, "work item type", "type", "ado work item type")
    if raw:
        return WORK_ITEM_TYPE_ALIASES.get(normalize(raw), str(raw).strip())
    file_name = normalize(row.get("_source_file"))
    if "deliver" in file_name:
        return "Deliverable"
    if "objective" in file_name:
        return "Objective"
    return "System process"


def _external_key(lookup: dict[str, Any], row: dict[str, Any]) -> str:
    value = _first_value(
        lookup,
        "process sequence id",
        "microsoft id",
        "partner id",
        "mavim id",
        "id",
        "external id",
        "key",
    )
    if value:
        return str(value).strip()
    stem = Path(str(row.get("_source_file") or "source")).stem
    return f"{stem}:{row.get('_source_sheet') or ''}:{row.get('_source_row')}"


def _parent_key(lookup: dict[str, Any], hierarchy_level: int | None) -> str | None:
    explicit = _first_value(
        lookup,
        "parent key",
        "parent id",
        "parent microsoft id",
        "parent process sequence id",
    )
    if explicit:
        return str(explicit).strip()
    return _derive_parent_from_sequence(_first_value(lookup, "process sequence id", "sequence id"))


def _derive_parent_from_sequence(value: Any) -> str | None:
    if not value:
        return None
    text = str(value).strip()
    dotted = text.split(".")
    if len(dotted) > 1 and all(part.isdigit() for part in dotted):
        return ".".join(dotted[:-1])
    for pattern in (r"^(.*)[.](\d+)$", r"^(.*)[/](.+)$", r"^(.*)[-](\d+)$"):
        match = re.match(pattern, text)
        if match:
            return match.group(1)
    return None


def _first_value(lookup: dict[str, Any], *keys: str) -> str | None:
    for key in keys:
        value = lookup.get(normalize(key))
        if value not in (None, ""):
            return str(value).strip()
    return None
